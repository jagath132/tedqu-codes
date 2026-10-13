"""
================================================================================
  TEACHER vs STUDENT DUPLICATE FINDER
  Checks for duplicate Email IDs and Phone Numbers across both files
  and creates a formatted Excel report.

  INSTALLATION:
      pip install pandas openpyxl

  HOW TO RUN:
      python find_teacher_student_duplicates.py

  WHAT IT DOES:
      1. Reads the Teacher Excel file
      2. Reads the Student Excel file
      3. Compares Email IDs and Phone Numbers between both files
      4. Creates a new Excel report showing all matching duplicates
================================================================================
"""

# ── Standard library ──────────────────────────────────────────────────────────
import os
import re

# ── Third-party libraries (install with: pip install pandas openpyxl) ─────────
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


# ==============================================================================
#  ██╗   ██╗███████╗███████╗██████╗     ███████╗███████╗████████╗████████╗██╗███╗   ██╗ ██████╗ ███████╗
#  ██║   ██║██╔════╝██╔════╝██╔══██╗    ██╔════╝██╔════╝╚══██╔══╝╚══██╔══╝██║████╗  ██║██╔════╝ ██╔════╝
#  ██║   ██║███████╗█████╗  ██████╔╝    ███████╗█████╗     ██║      ██║   ██║██╔██╗ ██║██║  ███╗███████╗
#  ██║   ██║╚════██║██╔══╝  ██╔══██╗    ╚════██║██╔══╝     ██║      ██║   ██║██║╚██╗██║██║   ██║╚════██║
#  ╚██████╔╝███████║███████╗██║  ██║    ███████║███████╗   ██║      ██║   ██║██║ ╚████║╚██████╔╝███████║
#   ╚═════╝ ╚══════╝╚══════╝╚═╝  ╚═╝    ╚══════╝╚══════╝   ╚═╝      ╚═╝   ╚═╝╚═╝  ╚═══╝ ╚═════╝ ╚══════╝
#
#  >>>  UPDATE ONLY THE VALUES BELOW — DO NOT CHANGE ANYTHING ELSE  <<<
# ==============================================================================

# ------------------------------------------------------------------------------
# FILE PATHS
#   → Change these to point to YOUR Excel files on your computer.
#   → Use forward slashes (/) or double back-slashes (\\) in the path.
#   → Example (Windows):  "C:/Users/YourName/Desktop/Teachers_Data.xlsx"
#   → Example (Mac/Linux): "/home/yourname/documents/Teachers_Data.xlsx"
# ------------------------------------------------------------------------------
TEACHER_FILE   = r"D:\daps  school\APS SCHOOL\teacher details\cleaned_teachers.xlsx"
# ⚠️  IMPORTANT — UPDATE THIS PATH  ⚠️
# This must point to your STUDENT ROSTER file (the one with the columns:
#   phone_number, email_id, student_name, academic_year, grade, standard,
#   parent_firstname, parent_lastname, date_of_birth, gender(M/F/O)
# It is NOT the "changes_report.xlsx" file — that file is a change-log
# and does not contain email_id / phone_number columns directly.
STUDENT_FILE   = r"D:\daps  school\APS SCHOOL\std details\cleaned_students.xlsx"
# ------------------------------------------------------------------------------
# OUTPUT FILE
#   → This MUST be a full FILE path ending in .xlsx — NOT just a folder.
#   → If you only give a folder here, the script will automatically save
#     the report as "duplicate_report.xlsx" inside that folder.
# ------------------------------------------------------------------------------
OUTPUT_FILE    = r"D:\daps  school\APS SCHOOL\std+teacher datat"
# ------------------------------------------------------------------------------
# SHEET NAMES
#   → The name of the sheet (tab) inside each Excel file.
#   → If you only have one sheet, leave these as None and
#     the script will automatically use the first sheet.
# ------------------------------------------------------------------------------
TEACHER_SHEET  = None   # ← e.g. "Sheet1" or None for first sheet
STUDENT_SHEET  = None   # ← e.g. "Sheet1" or None for first sheet

# ------------------------------------------------------------------------------
# COLUMN NAMES — TEACHER FILE
#   → These must EXACTLY match the column headers in the Teacher Excel file.
#   → They are case-sensitive.  "Email_ID" ≠ "email_id"
# ------------------------------------------------------------------------------
TEACHER_EMAIL_COL       = "email_id"        # ← teacher email column header
TEACHER_PHONE_COL       = "phone_number"    # ← teacher phone column header
TEACHER_FIRSTNAME_COL   = "firstname"       # ← teacher first name column header
TEACHER_LASTNAME_COL    = "lastname"        # ← teacher last name column header

# ------------------------------------------------------------------------------
# COLUMN NAMES — STUDENT FILE
#   → These must EXACTLY match the column headers in the Student Excel file.
#   → Based on the columns you shared, these are already correct:
#     phone_number, email_id, student_name, academic_year, grade, standard,
#     parent_firstname, parent_lastname, date_of_birth, gender(M/F/O)
# ------------------------------------------------------------------------------
STUDENT_EMAIL_COL  = "email_id"        # ← student email column header
STUDENT_PHONE_COL  = "phone_number"    # ← student phone column header
STUDENT_NAME_COL   = "student_name"    # ← student name column header

# ------------------------------------------------------------------------------
# FUTURE DUPLICATE-CHECK CONDITIONS
#   → If you want to check additional fields in the future (e.g. Date of Birth),
#     add your logic inside the function called  find_duplicates()  below,
#     following the same pattern used for email and phone.
# ------------------------------------------------------------------------------

# ==============================================================================
#  END OF USER SETTINGS — DO NOT CHANGE ANYTHING BELOW THIS LINE
# ==============================================================================


# ──────────────────────────────────────────────────────────────────────────────
#  HELPER FUNCTIONS
# ──────────────────────────────────────────────────────────────────────────────

def clean_email(value):
    """
    Cleans an email address so comparisons are consistent.
    Steps:
      - Convert to text (in case the value is a number)
      - Strip leading/trailing spaces
      - Convert to lowercase
      - Return empty string if the result is blank or 'nan'
    """
    if pd.isna(value):
        return ""
    val = str(value).strip().lower()
    return "" if val in ("nan", "none", "") else val


def clean_phone(value):
    """
    Cleans a phone number so comparisons are consistent.
    Steps:
      - Convert to text
      - Remove all non-digit characters (spaces, hyphens, +, brackets, etc.)
      - Return empty string if nothing is left
    """
    if pd.isna(value):
        return ""
    val = str(value).strip()
    digits_only = re.sub(r"\D", "", val)   # keep digits only
    return digits_only if digits_only else ""


def load_excel(filepath, sheet_name, label):
    """
    Loads an Excel file into a pandas DataFrame.
    - filepath   : path to the .xlsx file
    - sheet_name : sheet tab name (or None for the first sheet)
    - label      : friendly name used in error messages ("Teacher" or "Student")

    Returns the DataFrame, or stops the program with a clear error message.
    """
    if not os.path.exists(filepath):
        print(f"\n❌  ERROR — {label} file not found!")
        print(f"    Expected location : {os.path.abspath(filepath)}")
        print(f"    Please open this script and update the {label.upper()}_FILE setting.\n")
        raise SystemExit(1)

    try:
        # If no sheet name given, read the first sheet by using index 0
        sheet = sheet_name if sheet_name is not None else 0
        df = pd.read_excel(filepath, sheet_name=sheet)
        print(f"✅  Loaded {label} file  →  {len(df)} rows  |  {filepath}")
        return df
    except Exception as err:
        print(f"\n❌  ERROR — Could not read {label} file: {err}\n")
        raise SystemExit(1)


def check_required_columns(df, required_cols, label):
    """
    Verifies that all required column names exist in the DataFrame.
    - df            : the DataFrame to check
    - required_cols : list of column names that must be present
    - label         : "Teacher" or "Student" (used in error messages)

    Stops the program if any column is missing and tells the user which ones.
    """
    missing = [col for col in required_cols if col not in df.columns]
    if missing:
        print(f"\n❌  ERROR — Missing columns in {label} file:")
        for col in missing:
            print(f"      • '{col}'")
        print(f"\n    Available columns in {label} file:")
        for col in df.columns:
            print(f"      • '{col}'")
        print(f"\n    Please update the {label.upper()} column settings at the top of this script.\n")
        raise SystemExit(1)


def find_duplicates(teachers_df, students_df):
    """
    Core logic: compares teacher and student data to find shared
    Email IDs and/or Phone Numbers.

    Returns a list of dictionaries — one dictionary per duplicate match.
    Each dictionary becomes one row in the output Excel report.

    ─────────────────────────────────────────────────────────────────────
    TO ADD FUTURE DUPLICATE-CHECK CONDITIONS:
      Follow the same pattern used below for email and phone.
      For example, to also check Date of Birth:
        teacher_dob = str(teacher_row["date_of_birth"]).strip()
        student_dob = str(student_row["date_of_birth"]).strip()
        if teacher_dob and student_dob and teacher_dob == student_dob:
            ... (same record-building pattern)
    ─────────────────────────────────────────────────────────────────────
    """

    # ── Step 1: Build teacher name from first + last name columns ────────────
    teachers_df["_teacher_name"] = (
        teachers_df[TEACHER_FIRSTNAME_COL].fillna("").astype(str).str.strip()
        + " "
        + teachers_df[TEACHER_LASTNAME_COL].fillna("").astype(str).str.strip()
    ).str.strip()

    # ── Step 2: Clean email and phone columns in both DataFrames ─────────────
    teachers_df["_email_clean"] = teachers_df[TEACHER_EMAIL_COL].apply(clean_email)
    teachers_df["_phone_clean"] = teachers_df[TEACHER_PHONE_COL].apply(clean_phone)
    students_df["_email_clean"] = students_df[STUDENT_EMAIL_COL].apply(clean_email)
    students_df["_phone_clean"] = students_df[STUDENT_PHONE_COL].apply(clean_phone)

    # ── Step 3: Build lookup sets from student data ───────────────────────────
    #    (used for fast membership checks — much faster than looping twice)
    student_emails = set(students_df["_email_clean"].loc[students_df["_email_clean"] != ""])
    student_phones = set(students_df["_phone_clean"].loc[students_df["_phone_clean"] != ""])

    duplicate_rows = []

    # ── Step 4: Loop through every teacher row ────────────────────────────────
    for _, teacher_row in teachers_df.iterrows():
        t_email = teacher_row["_email_clean"]
        t_phone = teacher_row["_phone_clean"]
        t_name  = teacher_row["_teacher_name"]
        t_email_display = teacher_row[TEACHER_EMAIL_COL] if not pd.isna(teacher_row[TEACHER_EMAIL_COL]) else ""
        t_phone_display = teacher_row[TEACHER_PHONE_COL] if not pd.isna(teacher_row[TEACHER_PHONE_COL]) else ""

        email_match = t_email and t_email in student_emails
        phone_match = t_phone and t_phone in student_phones

        if not (email_match or phone_match):
            continue   # no duplicate — skip this teacher

        # ── Step 5: Find the matching student row(s) ──────────────────────────
        # We look for students who match by email OR phone individually so that
        # every combination of matches is captured as its own row.

        matched_students = set()   # will hold index values to avoid duplicates

        if email_match:
            email_matches = students_df[students_df["_email_clean"] == t_email]
            for s_idx in email_matches.index:
                matched_students.add(s_idx)

        if phone_match:
            phone_matches = students_df[students_df["_phone_clean"] == t_phone]
            for s_idx in phone_matches.index:
                matched_students.add(s_idx)

        # ── Step 6: Build one output row per teacher-student match ────────────
        for s_idx in matched_students:
            student_row = students_df.loc[s_idx]
            s_email = student_row["_email_clean"]
            s_phone = student_row["_phone_clean"]

            # Determine what kind of duplicate this teacher-student pair is
            is_email_dup = (t_email and s_email and t_email == s_email)
            is_phone_dup = (t_phone and s_phone and t_phone == s_phone)

            if is_email_dup and is_phone_dup:
                dup_type     = "Duplicate Email + Phone"
                matched_field = "Email ID + Phone Number"
                matched_value = f"Email: {t_email}  |  Phone: {t_phone}"
            elif is_email_dup:
                dup_type     = "Duplicate Email"
                matched_field = "Email ID"
                matched_value = t_email
            elif is_phone_dup:
                dup_type     = "Duplicate Phone"
                matched_field = "Phone Number"
                matched_value = t_phone
            else:
                continue   # shouldn't happen, but guard anyway

            s_name  = student_row[STUDENT_NAME_COL] if not pd.isna(student_row[STUDENT_NAME_COL]) else ""
            s_email_display = student_row[STUDENT_EMAIL_COL] if not pd.isna(student_row[STUDENT_EMAIL_COL]) else ""
            s_phone_display = student_row[STUDENT_PHONE_COL] if not pd.isna(student_row[STUDENT_PHONE_COL]) else ""

            remarks = "Teacher and Student duplicate found"

            duplicate_rows.append({
                "Duplicate Type"       : dup_type,
                "Matched Field"        : matched_field,
                "Matched Value"        : matched_value,
                "Teacher Name"         : t_name,
                "Teacher Email ID"     : t_email_display,
                "Teacher Phone Number" : t_phone_display,
                "Student Name"         : s_name,
                "Student Email ID"     : s_email_display,
                "Student Phone Number" : s_phone_display,
                "Remarks"              : remarks,
            })

    return duplicate_rows


def apply_excel_formatting(output_path, has_data):
    """
    Opens the freshly written Excel file and applies professional formatting:
      • Bold, coloured header row
      • Autofilter on headers
      • Auto-fit column widths
      • Centre-aligned key columns
      • Alternating row shading for readability
      • Thin borders on all cells
    """
    wb = load_workbook(output_path)
    ws = wb.active
    if ws is None:
        # wb.active is typed as Worksheet | None by openpyxl's stubs, but a
        # freshly-loaded workbook always has an active sheet. This assertion
        # narrows the type for Pylance and gives a clear error in the
        # (practically impossible) case it's ever actually missing.
        raise RuntimeError(f"Workbook '{output_path}' has no active worksheet.")

    # ── Colour palette ────────────────────────────────────────────────────────
    HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")   # dark navy
    ALT_ROW_FILL  = PatternFill("solid", fgColor="EBF3FB")   # light blue
    EMAIL_FILL    = PatternFill("solid", fgColor="FFF2CC")   # pale yellow
    PHONE_FILL    = PatternFill("solid", fgColor="E2EFDA")   # pale green
    BOTH_FILL     = PatternFill("solid", fgColor="FCE4D6")   # pale orange

    thin_side  = Side(style="thin", color="BFBFBF")
    thin_border = Border(left=thin_side, right=thin_side,
                         top=thin_side,  bottom=thin_side)

    # ── Format header row ─────────────────────────────────────────────────────
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for cell in ws[1]:
        cell.font      = header_font
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border    = thin_border
    ws.row_dimensions[1].height = 30

    # ── Format data rows ──────────────────────────────────────────────────────
    centre_cols = {1, 2, 3}   # Duplicate Type, Matched Field, Matched Value
    dup_type_col_idx = 1       # "Duplicate Type" is column A (index 1)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        # Choose row fill based on duplicate type
        # (cell.value is typed as a big union by openpyxl's stubs, so cast
        # explicitly to str before using the "in" operator on it)
        raw_dup_type_val = ws.cell(row=row_idx, column=dup_type_col_idx).value
        dup_type_val = str(raw_dup_type_val) if raw_dup_type_val is not None else ""
        if "Email + Phone" in dup_type_val:
            row_fill = BOTH_FILL
        elif "Email" in dup_type_val:
            row_fill = EMAIL_FILL
        elif "Phone" in dup_type_val:
            row_fill = PHONE_FILL
        else:
            row_fill = ALT_ROW_FILL if row_idx % 2 == 0 else PatternFill()

        for col_idx, cell in enumerate(row, start=1):
            cell.fill   = row_fill
            cell.border = thin_border
            if col_idx in centre_cols:
                cell.alignment = Alignment(horizontal="center",
                                           vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left",
                                           vertical="center", wrap_text=True)

    # ── Auto-fit column widths ────────────────────────────────────────────────
    for col_cells in ws.columns:
        max_len = 0
        col_index = col_cells[0].column
        if col_index is None:
            continue  # shouldn't happen for a real worksheet column, but guard for Pylance
        col_letter = get_column_letter(col_index)
        for cell in col_cells:
            try:
                cell_len = len(str(cell.value)) if cell.value else 0
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        adjusted_width = min(max_len + 4, 50)   # cap at 50 characters
        ws.column_dimensions[col_letter].width = adjusted_width

    # ── Add autofilter on header row ──────────────────────────────────────────
    ws.auto_filter.ref = ws.dimensions

    # ── Freeze the header row so it stays visible when scrolling ─────────────
    ws.freeze_panes = "A2"

    wb.save(output_path)


def resolve_output_path(output_path, default_filename="duplicate_report.xlsx"):
    """
    Makes sure we always end up with a proper FILE path ending in .xlsx.

    This is what fixes the crash you hit:
        ValueError: No engine for filetype: ''

    That error happens when OUTPUT_FILE points at a FOLDER (no filename,
    no extension) instead of a file. pandas can't figure out which Excel
    engine to use without a file extension.

    Behaviour:
      - If output_path already ends in .xlsx / .xls  → use it as-is.
      - If output_path is (or looks like) a folder    → append the
        default filename to it automatically.
    """
    root, ext = os.path.splitext(output_path)

    if ext.lower() in (".xlsx", ".xls"):
        return output_path

    # No recognizable file extension → treat it as a folder path
    print(f"\n⚠️   OUTPUT_FILE has no filename/extension — treating it as a folder.")
    print(f"    Will save the report as '{default_filename}' inside that folder.")
    return os.path.join(output_path, default_filename)


def save_report(duplicate_rows, output_path):
    """
    Saves the list of duplicate records to an Excel file.
    If no duplicates were found, creates the file with headers and
    one informational row.
    """
    # ── Make sure output_path is a proper FILE path (fixes folder-only paths) ─
    output_path = resolve_output_path(output_path)

    # ── Make sure the output folder exists ───────────────────────────────────
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)

    OUTPUT_COLUMNS = [
        "Duplicate Type",
        "Matched Field",
        "Matched Value",
        "Teacher Name",
        "Teacher Email ID",
        "Teacher Phone Number",
        "Student Name",
        "Student Email ID",
        "Student Phone Number",
        "Remarks",
    ]

    if duplicate_rows:
        df_out = pd.DataFrame(duplicate_rows, columns=OUTPUT_COLUMNS)
    else:
        # No duplicates found — write a single informational row
        df_out = pd.DataFrame([{
            "Duplicate Type"       : "No duplicates found",
            "Matched Field"        : "-",
            "Matched Value"        : "-",
            "Teacher Name"         : "-",
            "Teacher Email ID"     : "-",
            "Teacher Phone Number" : "-",
            "Student Name"         : "-",
            "Student Email ID"     : "-",
            "Student Phone Number" : "-",
            "Remarks"              : "No duplicates found between teacher and student data.",
        }], columns=OUTPUT_COLUMNS)

    df_out.to_excel(output_path, index=False)

    # Apply professional formatting
    apply_excel_formatting(output_path, has_data=bool(duplicate_rows))

    print(f"\n📄  Report saved  →  {os.path.abspath(output_path)}")
    print(f"    Total duplicate records written : {len(df_out)}")


# ──────────────────────────────────────────────────────────────────────────────
#  MAIN PROGRAM
# ──────────────────────────────────────────────────────────────────────────────

def main():
    print("\n" + "=" * 65)
    print("   TEACHER vs STUDENT DUPLICATE FINDER")
    print("=" * 65)

    # ── 1. Load both Excel files ──────────────────────────────────────────────
    teachers_df = load_excel(TEACHER_FILE, TEACHER_SHEET, "Teacher")
    students_df = load_excel(STUDENT_FILE, STUDENT_SHEET, "Student")

    # ── 2. Verify required columns exist ─────────────────────────────────────
    check_required_columns(
        teachers_df,
        [TEACHER_EMAIL_COL, TEACHER_PHONE_COL,
         TEACHER_FIRSTNAME_COL, TEACHER_LASTNAME_COL],
        "Teacher"
    )
    check_required_columns(
        students_df,
        [STUDENT_EMAIL_COL, STUDENT_PHONE_COL, STUDENT_NAME_COL],
        "Student"
    )

    print("\n🔍  Scanning for duplicate Email IDs and Phone Numbers ...")

    # ── 3. Find all duplicates ────────────────────────────────────────────────
    duplicate_rows = find_duplicates(teachers_df, students_df)

    # ── 4. Report findings to the console ────────────────────────────────────
    if duplicate_rows:
        print(f"\n⚠️   Found {len(duplicate_rows)} duplicate match(es)!")
        email_count = sum(1 for r in duplicate_rows if r["Duplicate Type"] == "Duplicate Email")
        phone_count = sum(1 for r in duplicate_rows if r["Duplicate Type"] == "Duplicate Phone")
        both_count  = sum(1 for r in duplicate_rows if "+" in r["Duplicate Type"])
        print(f"      • Duplicate Email only      : {email_count}")
        print(f"      • Duplicate Phone only      : {phone_count}")
        print(f"      • Duplicate Email + Phone   : {both_count}")
    else:
        print("\n✅  No duplicates found between teacher and student data.")

    # ── 5. Save the report ────────────────────────────────────────────────────
    save_report(duplicate_rows, OUTPUT_FILE)

    print("\n" + "=" * 65)
    print("   Done! Open the report file to see the results.")
    print("=" * 65 + "\n")


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    main()